import math

import openpyxl
import pytest

from escape_lab.activation_data import (
    ACTIVATION_HEADER,
    activation_records_from_rows,
    load_development_activation,
)
from escape_lab.mechanistic_pilot import activation_from_rows


def candidate(position=1, split="development", gene="FGFR1"):
    return {"candidate_id": f"{gene}:p.Ala{position}Gly", "gene": gene,
            "wt": "A", "mut": "G", "position": position, "split": split}


def row(nt="G1A", values=(1, 2, 3, 4), aa="A1G", gene="FGFR1"):
    return [gene, 1, "transcript", nt, aa, *values]


def test_distinct_codons_collapse_within_each_replicate_and_match_v1():
    candidates = [candidate()]
    rows = [row(values=(1, 2, 3, 4)), row("G1C", (3, 4, 5, 6))]
    records, audit = activation_records_from_rows(rows, candidates)
    record = records[candidates[0]["candidate_id"]]
    assert record == {"value": math.log2(4.5), "replicates": [2, 3, 4, 5], "active": True}
    old_values, old_audit = activation_from_rows(rows, candidates)
    assert {cid: r["value"] for cid, r in records.items()} == old_values
    assert all(audit[key] == value for key, value in old_audit.items())
    assert audit["eligible_nucleotide_rows"] == 2


@pytest.mark.parametrize("values,active", [
    ((1.5, 1.5, 1.5, 0), True), ((1.5, 1.5, 1.49, 8), True),
    ((1.5, 1.5, 1.49, 1.49), False), ((0, 0, 0, 0), False),
])
def test_active_requires_three_replicates_at_least_one_point_five(values, active):
    records, _ = activation_records_from_rows([row(values=values)], [candidate()])
    assert records[candidate()["candidate_id"]]["active"] is active


def test_missing_values_are_not_zero_and_complementary_codons_can_complete_replicates():
    records, audit = activation_records_from_rows(
        [row(values=(None, "NA", 2, 3)), row("G1C", (4, 5, None, "nan"))], [candidate()])
    assert records[candidate()["candidate_id"]]["replicates"] == [4, 5, 2, 3]
    assert audit["incomplete_protein_variants"] == 0
    records, audit = activation_records_from_rows([row(values=(None, 1, 1, 1))], [candidate()])
    assert records == {}
    assert audit["incomplete_protein_variants"] == 1


def test_excluded_id_skipped_before_any_measurement_or_nucleotide_access():
    class ExcludedRow:
        def __getitem__(self, index):
            if index == 0:
                return "FGFR1"
            if index == 4:
                return "A99G"
            raise AssertionError("Excluded row payload accessed")

    records, _ = activation_records_from_rows([ExcludedRow(), row()], [candidate()])
    assert set(records) == {candidate()["candidate_id"]}


@pytest.mark.parametrize("bad", [-1, float("inf"), float("-inf")])
def test_invalid_numeric_activation_rejected(bad):
    with pytest.raises(ValueError, match="finite nonnegative"):
        activation_records_from_rows([row(values=(bad, 1, 1, 1))], [candidate()])


def test_duplicate_nucleotide_identity_rejected_but_cross_gene_identity_allowed():
    with pytest.raises(ValueError, match="Duplicate nucleotide"):
        activation_records_from_rows([row(), row()], [candidate()])
    records, _ = activation_records_from_rows(
        [row(), row(gene="FGFR2")], [candidate(), candidate(gene="FGFR2")])
    assert len(records) == 2


def test_final_candidates_rejected_before_rows_are_read():
    class UnreadableRows:
        def __iter__(self):
            raise AssertionError("Rows accessed before split validation")

    with pytest.raises(ValueError, match="development"):
        activation_records_from_rows(UnreadableRows(), [candidate(split="final")])


def test_duplicate_candidate_identity_rejected():
    with pytest.raises(ValueError, match="Duplicate candidate"):
        activation_records_from_rows([], [candidate(), candidate()])


def test_loader_exact_headers_and_activation_columns_only(monkeypatch):
    calls = []

    class Sheet:
        def iter_rows(self, **kwargs):
            calls.append(kwargs)
            assert kwargs["max_col"] == 9
            assert kwargs["values_only"] is True
            return iter([ACTIVATION_HEADER] if kwargs["min_row"] == 10 else [row()])

    class Workbook:
        closed = False

        def __getitem__(self, name):
            assert name == "S10_Screen data table"
            return Sheet()

        def close(self):
            self.closed = True

    workbook = Workbook()

    def load(path, **kwargs):
        assert kwargs == {"read_only": True, "data_only": True}
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    records, _ = load_development_activation("unused.xlsx", [candidate()])
    assert records[candidate()["candidate_id"]]["replicates"] == [1, 2, 3, 4]
    assert workbook.closed
    assert calls == [{"min_row": 10, "max_row": 10, "max_col": 9, "values_only": True},
                     {"min_row": 11, "max_col": 9, "values_only": True}]


def test_loader_rejects_changed_header_and_closes_workbook(monkeypatch):
    class Workbook:
        closed = False

        def __getitem__(self, name):
            return self

        def iter_rows(self, **kwargs):
            return iter([(*ACTIVATION_HEADER[:5], "Resistance BR1", *ACTIVATION_HEADER[6:])])

        def close(self):
            self.closed = True

    workbook = Workbook()
    def load(*args, **kwargs):
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    with pytest.raises(ValueError, match="activation source columns"):
        load_development_activation("unused.xlsx", [candidate()])
    assert workbook.closed


def test_loader_rejects_final_candidates_before_opening_workbook(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Workbook opened before split validation")

    monkeypatch.setattr(openpyxl, "load_workbook", forbidden)
    with pytest.raises(ValueError, match="development"):
        load_development_activation("unused.xlsx", [candidate(split="final")])
