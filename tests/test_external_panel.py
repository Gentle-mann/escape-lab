import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from escape_lab import external_panel


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value))


@pytest.fixture
def authorized_panel(tmp_path):
    workbook = tmp_path / "panel.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Table S13"
    sheet["A8"] = "Compound"
    sheet["A9"] = "Futibatinib"
    sheet["A10"] = "Pemigatinib"
    variants = []
    for index in range(6):
        column = openpyxl.utils.get_column_letter(index + 2)
        variant = f"FGFR2:A{500 + index}V"
        label = f"FGFR2(A{500 + index}V)"
        sheet[f"{column}8"] = label
        sheet[f"{column}9"] = index * 10
        sheet[f"{column}10"] = index * 12
        variants.append({"variant_id": variant, "source_label": label,
                         "column": column, "group_id": f"block-{index // 2}"})
    book.save(workbook)
    manifest = tmp_path / "manifest.json"
    write_json(manifest, {"schema_version": "s13-external-panel-v1",
                         "panel_id": "synthetic", "document_path": str(workbook),
                         "document_sha256": digest(workbook), "sheet": "Table S13",
                         "header_row": 8, "drug_rows": {"fut": 9, "pem": 10},
                         "drug_labels": {"fut": "Futibatinib", "pem": "Pemigatinib"},
                         "variant_count": 6, "variants": variants})
    records = []
    for method in ["v1_fixed", "agent"]:
        for index, variant in enumerate(variants):
            for drug in ["pem", "fut"]:
                score = index / 6 if method == "agent" else 1 - index / 6
                records.append({"method": method, "variant_id": variant["variant_id"],
                                "drug": drug, "score": score})
    predictions = tmp_path / "predictions.json"
    write_json(predictions, {"predictions": records})
    freeze = tmp_path / "freeze.json"
    receipt = tmp_path / "external-receipt.json"
    contract = {"authorized": True, "panel_manifest_sha256": digest(manifest),
                "document_sha256": digest(workbook), "predictions_sha256": digest(predictions),
                "evaluator_sha256": digest(external_panel.__file__),
                "methods": ["v1_fixed", "agent"],
                "primary_metric": "mean_per_drug_spearman", "minimum_n_per_drug": 5}
    write_json(freeze, {"external_evaluation": contract})
    return {"manifest": manifest, "freeze": freeze, "predictions": predictions,
            "receipt": receipt, "workbook": workbook, "contract": contract}


def evaluate(files):
    return external_panel.evaluate_external_panel(
        files["manifest"], files["freeze"], files["predictions"], files["receipt"]
    )


def refreeze(files):
    manifest = json.loads(files["manifest"].read_text())
    manifest["document_sha256"] = digest(files["workbook"])
    write_json(files["manifest"], manifest)
    contract = files["contract"]
    contract["document_sha256"] = digest(files["workbook"])
    contract["panel_manifest_sha256"] = digest(files["manifest"])
    contract["predictions_sha256"] = digest(files["predictions"])
    write_json(files["freeze"], {"external_evaluation": contract})


@pytest.mark.parametrize("value, status, expected", [
    (0, "observed", 0.0), (15.2, "observed", 15.2),
    ("1.2e1", "observed", 12.0), ("2.5 × 10^1", "observed", 25.0),
    (None, "missing", None), ("NA", "missing", None),
    (">100", "censored_lower_bound", None), ("≤ 2.5", "censored_upper_bound", None),
    (">= 1e2", "censored_lower_bound", None), ("<= 2", "censored_upper_bound", None),
    (True, "invalid", None), ("not a measurement", "invalid", None),
    (float("inf"), "invalid", None), (float("nan"), "missing", None),
])
def test_measurement_status_preserves_missing_and_censored(value, status, expected):
    parsed = external_panel.parse_residual_activity(value)
    assert parsed["status"] == status
    assert parsed["value"] == expected


def test_out_of_range_measurements_are_flagged_without_clipping():
    parsed = external_panel.parse_residual_activity(-2)
    assert parsed == {"status": "observed", "value": -2.0,
                      "outside_nominal_range": True}


def test_metadata_inspection_does_not_return_outcomes(authorized_panel):
    metadata = external_panel.inspect_external_panel(authorized_panel["manifest"])
    assert metadata["variant_count"] == 6
    assert metadata["numeric_outcomes_inspected"] is False
    assert "value" not in metadata and "measurements" not in metadata


def test_complete_gate_and_metrics_and_global_receipt(authorized_panel):
    result = evaluate(authorized_panel)
    assert result["status"] == "completed"
    assert result["methods"]["agent"]["macro_spearman"] == pytest.approx(1)
    assert result["methods"]["v1_fixed"]["macro_spearman"] == pytest.approx(-1)
    assert "records" not in result
    assert "pvalue" not in json.dumps(result).lower()
    records = json.loads(Path(result["records_path"]).read_text())
    assert len(records["records"]) == 24
    assert records["records"][0]["measurement_status"] == "observed"
    receipt = json.loads(authorized_panel["receipt"].read_text())
    assert receipt["status"] == "completed"
    assert receipt["predictions_sha256"] == digest(authorized_panel["predictions"])
    original = authorized_panel["receipt"].read_bytes()
    with pytest.raises(FileExistsError):
        evaluate(authorized_panel)
    assert authorized_panel["receipt"].read_bytes() == original


def test_changed_predictions_rejected_before_label_read(authorized_panel, monkeypatch):
    authorized_panel["predictions"].write_text('{"predictions":[]}')
    def forbidden(*args, **kwargs):
        pytest.fail("Labels were accessed without valid prediction commitment")
    monkeypatch.setattr(external_panel, "_read_numeric_cells", forbidden)
    with pytest.raises(ValueError, match="prediction"):
        evaluate(authorized_panel)
    assert not authorized_panel["receipt"].exists()


def test_incomplete_prediction_coverage_rejected_before_reveal(authorized_panel, monkeypatch):
    payload = json.loads(authorized_panel["predictions"].read_text())
    payload["predictions"].pop()
    write_json(authorized_panel["predictions"], payload)
    refreeze(authorized_panel)
    def forbidden(*args, **kwargs):
        pytest.fail("Labels were accessed before complete predictions")
    monkeypatch.setattr(external_panel, "_read_numeric_cells", forbidden)
    with pytest.raises(ValueError, match="coverage"):
        evaluate(authorized_panel)


def test_unrecognized_or_duplicate_method_predictions_rejected(authorized_panel):
    payload = json.loads(authorized_panel["predictions"].read_text())
    payload["predictions"].append(payload["predictions"][0])
    write_json(authorized_panel["predictions"], payload)
    refreeze(authorized_panel)
    with pytest.raises(ValueError, match="Duplicate"):
        evaluate(authorized_panel)


def test_authorization_requires_explicit_true_and_evaluator_hash(authorized_panel):
    authorized_panel["contract"]["authorized"] = "true"
    refreeze(authorized_panel)
    with pytest.raises(ValueError, match="authoriz"):
        evaluate(authorized_panel)
    authorized_panel["contract"]["authorized"] = True
    authorized_panel["contract"]["evaluator_sha256"] = "0" * 64
    refreeze(authorized_panel)
    with pytest.raises(ValueError, match="evaluator"):
        evaluate(authorized_panel)


def test_wrong_header_spends_receipt_and_does_not_read_labels(authorized_panel, monkeypatch):
    book = openpyxl.load_workbook(authorized_panel["workbook"])
    book.active["B8"] = "FGFR2:A500W"
    book.save(authorized_panel["workbook"])
    refreeze(authorized_panel)
    def forbidden(*args, **kwargs):
        pytest.fail("Labels read despite metadata mismatch")
    monkeypatch.setattr(external_panel, "_read_numeric_cells", forbidden)
    with pytest.raises(ValueError, match="header"):
        evaluate(authorized_panel)
    assert json.loads(authorized_panel["receipt"].read_text())["status"] == "failed_spent"
    with pytest.raises(FileExistsError):
        evaluate(authorized_panel)


def test_missing_censored_and_formula_are_not_zero_imputed(authorized_panel):
    book = openpyxl.load_workbook(authorized_panel["workbook"])
    sheet = book.active
    sheet["B9"] = None
    sheet["C9"] = ">100"
    sheet["D9"] = "=1+1"
    book.save(authorized_panel["workbook"])
    refreeze(authorized_panel)
    result = evaluate(authorized_panel)
    assert result["measurement_coverage"]["fut"]["observed"] == 3
    assert result["methods"]["agent"]["per_drug"]["fut"]["n"] == 3
    assert result["methods"]["agent"]["macro_spearman"] is None
    assert result["methods"]["agent"]["per_drug"]["pem"]["spearman"] == pytest.approx(1)


def test_constant_prediction_is_not_fabricated_zero_correlation(authorized_panel):
    payload = json.loads(authorized_panel["predictions"].read_text())
    for row in payload["predictions"]:
        if row["method"] == "agent":
            row["score"] = 0.5
    write_json(authorized_panel["predictions"], payload)
    refreeze(authorized_panel)
    result = evaluate(authorized_panel)
    assert result["methods"]["agent"]["macro_spearman"] is None
    assert result["methods"]["agent"]["per_drug"]["pem"]["reason"] == "constant_predictions"


def test_read_failure_remains_spent(authorized_panel, monkeypatch):
    def fail(*args, **kwargs):
        assert json.loads(authorized_panel["receipt"].read_text())["status"] == "reserved"
        raise RuntimeError("synthetic read failure")
    monkeypatch.setattr(external_panel, "_read_numeric_cells", fail)
    with pytest.raises(RuntimeError, match="synthetic read failure"):
        evaluate(authorized_panel)
    assert json.loads(authorized_panel["receipt"].read_text())["status"] == "failed_spent"


def test_prediction_mutation_after_reservation_prevents_label_access(authorized_panel, monkeypatch):
    original_inspect = external_panel.inspect_external_panel
    def tamper(path):
        result = original_inspect(path)
        authorized_panel["predictions"].write_text('{"predictions":[]}')
        return result
    def forbidden(*args, **kwargs):
        pytest.fail("Outcome cells read after prediction commitment changed")
    monkeypatch.setattr(external_panel, "inspect_external_panel", tamper)
    monkeypatch.setattr(external_panel, "_read_numeric_cells", forbidden)
    with pytest.raises(ValueError, match="predictions changed"):
        evaluate(authorized_panel)
    assert json.loads(authorized_panel["receipt"].read_text())["status"] == "failed_spent"


def test_private_reader_requires_the_exact_live_receipt(authorized_panel):
    authorization = external_panel._Authorization(authorized_panel["receipt"], "test", {})
    write_json(authorized_panel["receipt"], {"status": "completed", "authorization_id": "test",
                                           "bound_files": {}})
    with pytest.raises(ValueError, match="not authorized"):
        external_panel._read_numeric_cells(json.loads(authorized_panel["manifest"].read_text()),
                                          authorization)
