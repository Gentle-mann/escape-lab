import hashlib
import io
import json
import sys
import zipfile

import numpy as np
import openpyxl
import pytest

from escape_lab import reproduction


def test_frozen_recipe_is_development_only_and_grouped():
    recipe = reproduction.load_recipe()
    candidates = recipe["candidates"]
    assert len(candidates) == 668
    assert len({candidate["group_id"] for candidate in candidates}) == 10
    assert all(candidate["split"] in {"development", "validation"} for candidate in candidates)
    assert not ({candidate["group_id"] for candidate in candidates} & set(recipe["reserved_groups"]))
    assert set(candidates[0]) == set(reproduction.CANDIDATE_FIELDS)
    mutated = json.loads(json.dumps(recipe))
    mutated["candidates"][0]["split"] = "final"
    with pytest.raises(ValueError, match="identity digest"):
        reproduction.validate_recipe(mutated)


def test_predictor_export_allows_only_frozen_identities_and_collapses_codons(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S10_Screen data table"
    headers = [None] * 43
    headers[0], headers[4] = "FGFR", "AA MUT"
    headers[41], headers[42] = "AlphaMissense score", "EVE score"
    sheet.append([])
    for column, value in enumerate(headers, 1):
        sheet.cell(10, column, value)
    for row_index, (gene, aa, alpha, eve) in enumerate([
        ("FGFR2", "N549K", .2, .3), ("FGFR2", "N549K", .2, .6),
        ("FGFR2", "V564F", .4, "."), ("FGFR3", "N549K", "DO NOT PARSE", "NO"),
    ], 11):
        for column, value in ((1, gene), (5, aa), (42, alpha), (43, eve)):
            sheet.cell(row_index, column, value)
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    candidates = [
        {"candidate_id": "FGFR2:p.Asn549Lys", "gene": "FGFR2", "wt": "N", "position": 549,
         "mut": "K", "split": "development"},
        {"candidate_id": "FGFR2:p.Val564Phe", "gene": "FGFR2", "wt": "V", "position": 564,
         "mut": "F", "split": "validation"},
    ]
    rows = reproduction.predictor_rows(source, candidates)
    assert rows[0]["alphamissense"] == .2
    assert rows[0]["eve"] == ""
    assert rows[0]["eve_status"] == "discordant_not_collapsed"
    assert rows[1]["eve_status"] == "missing"
    assert len(rows) == 2


def test_predictor_export_rejects_final_before_opening_source(tmp_path):
    with pytest.raises(ValueError, match="development"):
        reproduction.predictor_rows(tmp_path / "absent.xlsx", [{"split": "final"}])


def test_download_digest_mismatch_never_installs_file(tmp_path):
    destination = tmp_path / "input.json"
    with pytest.raises(ValueError, match="checksum"):
        reproduction.write_verified(destination, b"changed source", "0" * 64)
    assert not destination.exists()
    payload = b"pinned source"
    reproduction.write_verified(destination, payload, hashlib.sha256(payload).hexdigest())
    assert destination.read_bytes() == payload
    with pytest.raises(FileExistsError):
        reproduction.write_verified(destination, payload, hashlib.sha256(payload).hexdigest())


def test_supplement_extraction_uses_exact_member_and_hash():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("other.xlsx", b"wrong")
        archive.writestr("expected.xlsx", b"correct")
    result = reproduction.supplement_bytes(output.getvalue(), "expected.xlsx",
                                          hashlib.sha256(b"correct").hexdigest())
    assert result == b"correct"
    with pytest.raises(ValueError, match="checksum"):
        reproduction.supplement_bytes(output.getvalue(), "other.xlsx", "0" * 64)


def test_synthetic_demo_exposes_rounding_without_scientific_claim():
    result = reproduction.synthetic_saturation_demo()
    assert result["data_kind"] == "synthetic"
    assert result["probability_ap"] == .5
    assert result["margin_ap"] == 1.
    assert result["fits"] == 0
    assert result["api_calls"] == 0


def test_doctor_never_downloads_or_reads_workbook(monkeypatch, tmp_path):
    def forbidden(*args, **kwargs):
        raise AssertionError("doctor must not access network or scientific outcomes")
    monkeypatch.setattr(reproduction, "fetch_bytes", forbidden)
    monkeypatch.setattr(reproduction, "load_development_activation", forbidden)
    result = reproduction.doctor(tmp_path)
    assert result["recipe"]["variants"] == 668
    assert not result["prepared_inputs_complete"]


def test_doctor_rejects_receipt_without_actual_inputs(tmp_path):
    receipt = {"status": "prepared", "recipe_sha256": reproduction.sha256(reproduction.RECIPE_PATH),
               "input_sha256": {}}
    (tmp_path / "prepared.json").write_text(json.dumps(receipt))
    result = reproduction.doctor(tmp_path)
    assert not result["prepared_inputs_complete"]
    assert "missing required inputs" in result["prepared_inputs_error"]


def test_full_run_refuses_existing_output_and_incomplete_preparation(tmp_path):
    output = tmp_path / "result"
    output.mkdir()
    with pytest.raises(FileExistsError):
        reproduction.reproduce_headline(tmp_path / "inputs", output)
    with pytest.raises(ValueError, match="prepare"):
        reproduction.reproduce_headline(tmp_path / "inputs", tmp_path / "fresh")


def test_summary_must_match_both_arms_and_hits():
    recipe = reproduction.load_recipe()
    expected = recipe["expected_headline"]
    actual = {arm: {"mean_fold_high_ap": values["mean_fold_high_ap"],
                    "high_hits": values["high_hits"], "budget": 68}
              for arm, values in expected.items()}
    assert reproduction.headline_comparison(actual, recipe)["matches_published_aggregate"]
    actual["margin"]["mean_fold_high_ap"] = np.nextafter(.1, 1.)
    assert not reproduction.headline_comparison(actual, recipe)["matches_published_aggregate"]


@pytest.mark.parametrize("matches,exit_code", [(True, 0), (False, 1)])
def test_headline_cli_returns_nonzero_when_aggregate_reproduction_fails(
    monkeypatch, capsys, matches, exit_code
):
    def simulated_run(*args, **kwargs):
        return {"status": "completed", "comparison": {"matches_published_aggregate": matches},
                "fit_count": 3, "fitting_seconds": 1., "elapsed_seconds": 2.,
                "feature_mode": "fresh extraction"}
    monkeypatch.setattr(reproduction, "reproduce_headline", simulated_run)
    monkeypatch.setattr(sys, "argv", ["escape-lab-reproduce", "headline", "--inputs", "unused",
                                      "--output", "unused-output"])
    assert reproduction.main() == exit_code
    assert json.loads(capsys.readouterr().out)["comparison"]["matches_published_aggregate"] is matches
