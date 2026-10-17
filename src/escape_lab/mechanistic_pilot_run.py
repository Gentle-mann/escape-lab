"""Freeze, run once, and report the small development-only activation pilot."""

import argparse
import json
import time
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import openpyxl
from sklearn.metrics import average_precision_score

from escape_lab.biological_features import BiologicalFeatureStore
from escape_lab.data import sha256
from escape_lab.improvement import load_labels, write_json
from escape_lab.mechanistic_pilot import (
    ACTIVATION_PARAMS,
    ARMS,
    SALT,
    SEED,
    activation_from_rows,
    activation_metrics,
    advancement_gate,
    eligible_candidates,
    evaluate_pilot,
)
from escape_lab.revision_evaluation import grouped_folds

WORKBOOK = Path("research/external-predictions/2026-10-17-aggregate-sources/raw/"
                "fgfr-supplementary-tables.xlsx")
MANO_AUDIT = Path("research/mechanistic-map/mano-metadata-audit.json")
RLY_AUDIT = Path("data/external_discovery/source_audit/external-alignment-metadata.json")


def protected_artifacts(root):
    paths = set()
    for directory in ("src", "tests", "data", "reports"):
        for p in (root / directory).rglob("*"):
            if p.is_file() and "__pycache__" not in p.parts and p.suffix != ".pyc":
                paths.add(p)
    paths.update(root / p for p in (WORKBOOK, MANO_AUDIT, RLY_AUDIT))
    return {str(p.relative_to(root)): sha256(p) for p in sorted(paths)}


def paired_intervals(records, replicates=1000):
    """Descriptive block bootstrap conditional on already fitted OOF predictions."""
    groups = sorted({r["group_id"] for r in records["structure"]})
    index = {g: i for i, g in enumerate(groups)}
    reference = {(r["candidate_id"], r["drug"]): r for r in records["structure"]}
    arrays = {}
    for arm in ARMS:
        rows = {(r["candidate_id"], r["drug"]): r for r in records[arm]}
        if set(rows) != set(reference):
            raise ValueError("Bootstrap requires matched predictions")
        if any(any(rows[k][field] != reference[k][field] for field in ("label", "group_id"))
               for k in reference):
            raise ValueError("Bootstrap methods disagree on labels/groups")
        arrays[arm] = {}
        for drug in ("pem", "fut"):
            keys = sorted(k for k in rows if k[1] == drug and type(rows[k]["label"]) is bool)
            arrays[arm][drug] = (
                np.asarray([rows[k]["label"] for k in keys], dtype=int),
                np.asarray([rows[k]["probability"] for k in keys]),
                np.asarray([index[rows[k]["group_id"]] for k in keys]),
            )
    rng = np.random.default_rng(SEED)
    differences = {arm: [] for arm in ARMS[1:]}
    zero_positive_draws = 0
    for _ in range(replicates):
        counts = np.bincount(rng.integers(len(groups), size=len(groups)), minlength=len(groups))
        scores, degenerate = {}, False
        for arm in ARMS:
            aps = []
            for y, p, group in arrays[arm].values():
                weights = counts[group]
                if np.dot(y, weights) == 0:
                    aps.append(0.0)
                    degenerate = True
                else:
                    aps.append(float(average_precision_score(y, p, sample_weight=weights)))
            scores[arm] = float(np.mean(aps))
        zero_positive_draws += int(degenerate)
        for arm, values in differences.items():
            values.append(scores[arm] - scores["structure"])
    return {"replicates": replicates, "seed": SEED, "aligned_groups": len(groups),
            "draws_with_a_zero_positive_drug": zero_positive_draws,
            "macro_ap_difference_vs_structure": {
                arm: {"percentile_2_5": float(np.percentile(values, 2.5)),
                      "percentile_97_5": float(np.percentile(values, 97.5))}
                for arm, values in differences.items()},
            "interpretation": "Descriptive only: 11 related blocks, no refitting, no biological "
                              "replicate resampling; not a generalization confidence guarantee."}


def report_text(result, run_dir):
    lines = ["# Activation-transfer viability pilot", "",
             f"Decision: **{result['gate']['status']}**. This is development triage, not a fresh test.",
             "", (f"Candidates: {result['cohort']['n']}; aligned blocks: "
             f"{result['cohort']['groups']}; total run: {result['elapsed_seconds']:.2f} seconds."),
             "", ("Four fixed arms use the same three outer group folds and resistance labels. "
             "Predicted/shuffled activation features are generated with nested group cross-fitting. "
             "Measured activation is an assay-assisted diagnostic, unavailable for sequence-only inference."),
             "", "| Arm | Pem AP | Fut AP | Macro AP | Macro Brier |",
             "|---|---:|---:|---:|---:|"]
    for arm in ARMS:
        m = result["metrics"][arm]
        lines.append(f"| {arm} | {m['per_drug']['pem']['average_precision']:.5f} | "
                     f"{m['per_drug']['fut']['average_precision']:.5f} | "
                     f"{m['macro_ap']:.5f} | {m['macro_brier']:.5f} |")
    lines.extend(["", ("AP is average precision, a ranking metric; it is not percentage accuracy. "
                  "Higher AP and lower Brier are better."), "", "## Sample size and fold results", ""])
    base = result["metrics"]["structure"]
    for drug in ("pem", "fut"):
        m = base["per_drug"][drug]
        lines.append(f"- {drug}: {m['n']} conclusive variants, {m['positive_count']} resistant, "
                     f"{m['inconclusive_count']} inconclusive labels; prevalence {m['prevalence']:.5f}.")
    lines.extend(["", ("| Fold | Baseline macro AP | Predicted activation | Measured activation | "
                  "Shuffled activation | Pem positives | Fut positives |"),
                  "|---|---:|---:|---:|---:|---:|---:|"])
    for fold in base["by_fold"]:
        vals = [result["metrics"][arm]["by_fold"][fold]["macro_ap"] for arm in ARMS]
        m = base["by_fold"][fold]["per_drug"]
        lines.append(f"| {fold} | " + " | ".join(f"{v:.5f}" for v in vals)
                     + f" | {m['pem']['positive_count']} | {m['fut']['positive_count']} |")
    lines.extend(["", "## Activation prediction", "",
                  ("Target: log2(1 + median activation enrichment). The four assay replicates "
                  "are summarized after collapsing nucleotide representations within each replicate."),
                  "", "| Arm / gene | Spearman | Mean absolute error | N |",
                  "|---|---:|---:|---:|"])
    for arm, metrics in result["activation_metrics"].items():
        for name, m in {"pooled": metrics["pooled"], **metrics["per_gene"]}.items():
            rho = "undefined" if m["spearman"] is None else f"{m['spearman']:.4f}"
            lines.append(f"| {arm} / {name} | {rho} | {m['mae_log2_1p']:.4f} | {m['n']} |")
    lines.extend(["", "## Predeclared decision", "", "```json",
                  json.dumps(result["gate"], indent=2), "```", "", "## Descriptive uncertainty", ""])
    for arm, interval in result["uncertainty"]["macro_ap_difference_vs_structure"].items():
        lines.append(f"- {arm} minus structure, block-bootstrap 2.5th–97.5th percentiles: "
                     f"[{interval['percentile_2_5']:.5f}, {interval['percentile_97_5']:.5f}].")
    lines.extend(["", result["uncertainty"]["interpretation"], "", "## Scope and integrity", "",
                  ("- Original development and validation data have already informed prior work. "
                  "This run is exploratory, despite grouped cross-validation."),
                  "- The 11 blocks are related sequence regions, not 11 independent biological experiments.",
                  ("- Activation and resistance come from different cell backgrounds. Enrichment is "
                  "not an isolated kinase activity or absolute growth-rate measurement."),
                  ("- Structural distances are proxies, not simulated mutant binding affinities. "
                  "This pilot does not test kinetics, drug combinations, cell evolution or patient outcomes."),
                  ("- The predicted branch is nonlinear auxiliary supervision. Its features include "
                  "predicted activation and its interaction with fixed ligand proximity; no LLM is used."),
                  "- A single within-gene shuffle is a capacity control, not a permutation significance test.",
                  ("- Complete activation coverage defines the common cohort. Missing assay values "
                  "are not imputed as zero."),
                  ("- Original final and prospective RLY/MANO aligned blocks are excluded from "
                  "activation and resistance fits. External numerical tables are not read."),
                  ("- The trusted evaluator parses the monolithic original outcome JSON and source "
                  "workbook but exposes only authorized development rows/columns to this pilot."),
                  (f"- {result['integrity']['protected_file_count']} pre-existing source, test, data and "
                  "report artifacts remained hash-identical. Original evaluation receipts are unchanged."),
                  "", ("Protocol was written and hashed before numerical ingestion. No hyperparameter "
                  "search, LLM calls, external-test opening or post-score model revision occurred."),
                  "", f"Run directory: `{run_dir}`.", ""])
    return "\n".join(lines)


def run(root, output):
    root, output = Path(root).resolve(), Path(output).resolve()
    if output.exists():
        raise FileExistsError("Pilot output must be a fresh directory; preserve every attempt")
    start = time.monotonic()
    protected = protected_artifacts(root)
    candidates = json.loads((root / "data/processed/candidates.json").read_text())
    mano = json.loads((root / MANO_AUDIT).read_text())
    rly = json.loads((root / RLY_AUDIT).read_text())
    reserved = set(mano["mapped_groups"]) | {c["group_id"] for c in rly["panel_variants"]["rly2023"]}
    eligible = eligible_candidates(candidates, reserved)
    if len(eligible) != 2749 or len({c["group_id"] for c in eligible}) != 11:
        raise ValueError("Cohort differs from metadata-only research audit")
    tasks = grouped_folds(eligible, SALT, prefix="pilot")
    protocol = {
        "version": "activation-transfer-pilot-v1", "frozen_at": datetime.now(UTC).isoformat(),
        "hypothesis": "Activation-informed nonlinear features improve grouped resistance ranking.",
        "scope": "Previously used development+validation only; no fresh generalization claim",
        "eligible_candidate_ids": [c["candidate_id"] for c in eligible],
        "reserved_groups": sorted(reserved), "seed": SEED, "salt": SALT, "arms": ARMS,
        "fold_groups": {t["task_id"]: sorted({c["group_id"] for c in t["pool"]}) for t in tasks},
        "activation_model": ACTIVATION_PARAMS, "resistance_model": "structure-logistic-c0.1",
        "activation_features": "chemistry and intrinsic pem-template protein geometry; exclude "
                               "ligand and covalent_cys fields; gene indicators; train-only imputation",
        "augmentation": "activation + activation * exp(-ligand_CA_distance / 5), zero if unavailable",
        "activation_transform": "log2(1 + median of four per-replicate codon medians)",
        "missingness": "require four available per-protein replicate summaries; use common cohort",
        "shuffle": "within gene, separately inside every training fit, deterministic salt seed",
        "gate": {"macro_ap_gain_min": .005, "positive_fold_gains_min": 2,
                 "per_drug_ap_loss_max": .01, "per_drug_brier_max": "baseline * 1.05 + .001",
                 "must_beat_shuffled": True, "activation_spearman_must_be_positive": True,
                 "inconclusive_if": "zero-positive drug/fold or any fallback/convergence warning"},
        "uncertainty": "1000 paired aligned-block bootstrap draws, seed17, descriptive only",
        "protected_sha256": protected,
    }
    output.mkdir(parents=True)
    write_json(output / "protocol.json", protocol)
    protocol_hash = sha256(output / "protocol.json")
    write_json(output / "status.json", {"status": "started", "protocol_sha256": protocol_hash})
    try:
        workbook = openpyxl.load_workbook(root / WORKBOOK, read_only=True, data_only=True)
        try:
            sheet = workbook["S10_Screen data table"]
            header = next(sheet.iter_rows(min_row=10, max_row=10, max_col=9, values_only=True))
            if list(header) != ["FGFR", "Index", "Transcript ID", "NT MUT", "AA MUT",
                                "Activation BR1", "Activation BR2", "Activation BR3", "Activation BR4"]:
                raise ValueError("Unexpected activation source columns")
            activation, coverage = activation_from_rows(
                sheet.iter_rows(min_row=11, max_col=9, values_only=True), eligible)
        finally:
            workbook.close()
        cohort = [c for c in eligible if c["candidate_id"] in activation]
        if {c["group_id"] for c in cohort} != {c["group_id"] for c in eligible}:
            raise ValueError("Missing activation removed an entire predeclared fold group")
        labels = load_labels(root / "data/processed", cohort)
        store = BiologicalFeatureStore(root / "data/revision-v2/features")
        fitting_start = time.monotonic()
        raw_result = evaluate_pilot(cohort, labels, activation, store)
        fitting_seconds = time.monotonic() - fitting_start
        act_metrics = activation_metrics(raw_result["activation_records"])
        gate = advancement_gate(raw_result, act_metrics)
        intervals = paired_intervals(raw_result["records"])
        changed = [p for p, h in protected.items() if sha256(root / p) != h]
        if changed or sha256(output / "protocol.json") != protocol_hash:
            raise ValueError(f"Protected artifacts changed: {changed}")
        result = {"protocol_sha256": protocol_hash, "cohort": {
            "n": len(cohort), "groups": len({c["group_id"] for c in cohort}),
            "by_gene": dict(Counter(c["gene"] for c in cohort))},
            "coverage": coverage, "metrics": raw_result["metrics"],
            "activation_metrics": act_metrics, "gate": gate, "uncertainty": intervals,
            "model_diagnostics": raw_result["diagnostics"], "fitting_seconds": fitting_seconds,
            "elapsed_seconds": time.monotonic() - start,
            "integrity": {"protected_file_count": len(protected), "changed_files": changed}}
        # Row-level developmental outcomes are evaluator-owned; never print them or feed an agent.
        write_json(output / "private-predictions.json", raw_result)
        (output / "private-predictions.json").chmod(0o600)
        write_json(output / "results.json", result)
        (output / "report.md").write_text(report_text(result, output))
        write_json(output / "status.json", {"status": "completed", "protocol_sha256": protocol_hash,
                                           "results_sha256": sha256(output / "results.json")})
        return result
    except Exception as error:
        write_json(output / "status.json", {"status": "failed", "protocol_sha256": protocol_hash,
                                           "error_type": type(error).__name__, "error": str(error)})
        raise


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = run(args.root, args.output)
    print(json.dumps({"gate": result["gate"], "cohort": result["cohort"],
                      "fitting_seconds": result["fitting_seconds"],
                      "elapsed_seconds": result["elapsed_seconds"],
                      "output": str(args.output.resolve())}, indent=2))
